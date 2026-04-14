"""
Smart Exam Analyzer — Flask Server
Auto-detects subject from uploaded files and generates important questions via RAG + LLM.
"""

import os, sys, io
from collections import Counter
from flask import Flask, request, jsonify, send_from_directory

GROQ_API_KEY = os.getenv("GROQ_API_KEY")
GROQ_MODEL   = "llama-3.1-8b-instant"

if not GROQ_API_KEY:
    sys.exit("[ERROR] GROQ_API_KEY not found. Set it in your environment.")
# ── IMPORTS ─────────────────────────────────────────────────
try:
    import chromadb
    from sentence_transformers import SentenceTransformer
    from groq import Groq
except ImportError as e:
    sys.exit(f"[ERROR] Missing library: {e}\nRun: pip install -r requirements.txt")

# ── GLOBALS ──────────────────────────────────────────────────
embed_model = None
chroma_client = None

def init_db():
    global embed_model, chroma_client
    print("[INIT] Loading embedding model...")
    embed_model   = SentenceTransformer("sentence-transformers/all-MiniLM-L6-v2")
    chroma_client = chromadb.Client()
    print("[INIT] Ready.")

# ── FLASK APP ────────────────────────────────────────────────
app = Flask(__name__, static_folder=".")

@app.route("/")
def index():
    return send_from_directory(".", "exam_upload.html")

@app.route("/analyze", methods=["POST"])
def analyze():
    uploaded_files = request.files.getlist("files")
    if not uploaded_files:
        return jsonify({"error": "No files uploaded."}), 400

    # 1. Extract all text from uploaded files
    all_lines = []
    for f in uploaded_files:
        all_lines.extend(extract_lines(f))

    if not all_lines:
        return jsonify({"error": "Could not extract any text from the uploaded files."}), 400

    full_text = "\n".join(all_lines)

    # 2. Auto-detect subject using LLM
    subject = detect_subject(full_text)

    # 3. Extract topics from the content using LLM
    topics_raw = extract_topics(full_text, subject)

    # 4. Build ChromaDB from extracted lines
    collection = build_collection(all_lines)

    # 5. Generate important questions per topic via RAG
    important_questions = []
    for topic in topics_raw[:8]:  # top 8 topics
        q = generate_important_question(topic, collection, subject)
        important_questions.append({"topic": topic, "question": q})

    # 6. Generate a study priority summary
    summary = generate_summary(subject, topics_raw)

    return jsonify({
        "subject":   subject,
        "topics":    topics_raw,
        "questions": important_questions,
        "summary":   summary,
        "total_lines": len(all_lines),
    })

@app.route("/ask", methods=["POST"])
def ask():
    data  = request.get_json()
    query = data.get("query", "").strip()
    subject = data.get("subject", "")
    if not query:
        return jsonify({"error": "Empty query"}), 400
    answer = llm_ask(query, subject)
    return jsonify({"answer": answer})

# ── FILE PARSING ─────────────────────────────────────────────
def extract_lines(file_obj):
    name = file_obj.filename.lower()
    ext  = os.path.splitext(name)[1]
    lines = []
    try:
        if ext == ".pdf":
            import PyPDF2
            reader = PyPDF2.PdfReader(file_obj)
            for page in reader.pages:
                lines.extend((page.extract_text() or "").split("\n"))

        elif ext == ".docx":
            import docx
            doc = docx.Document(file_obj)
            for para in doc.paragraphs:
                lines.append(para.text)

        elif ext in (".xlsx", ".xls"):
            import openpyxl
            wb = openpyxl.load_workbook(file_obj, read_only=True, data_only=True)
            for ws in wb.worksheets:
                for row in ws.iter_rows(values_only=True):
                    lines.append(" ".join(str(c) for c in row if c))

        elif ext == ".csv":
            import csv
            content = file_obj.read().decode("utf-8", errors="ignore")
            for row in csv.reader(io.StringIO(content)):
                lines.append(" ".join(row))

        elif ext in (".png", ".jpg", ".jpeg", ".bmp", ".tiff", ".webp"):
            try:
                import pytesseract
                from PIL import Image
                lines.extend(pytesseract.image_to_string(Image.open(file_obj)).split("\n"))
            except ImportError:
                print("[WARN] Install pytesseract + Pillow for image OCR.")

        elif ext in (".txt", ".md", ".text"):
            lines.extend(file_obj.read().decode("utf-8", errors="ignore").split("\n"))

        elif ext == ".json":
            import json
            data = json.loads(file_obj.read().decode("utf-8", errors="ignore"))
            if isinstance(data, list):
                for item in data:
                    lines.append(item if isinstance(item, str) else item.get("question", str(item)))
        else:
            lines.extend(file_obj.read().decode("utf-8", errors="ignore").split("\n"))

    except Exception as e:
        print(f"[WARN] Could not parse {file_obj.filename}: {e}")

    cleaned = [l.strip() for l in lines if len(l.strip()) > 15]
    print(f"[PARSE] {file_obj.filename} → {len(cleaned)} lines")
    return cleaned

# ── CHROMADB ─────────────────────────────────────────────────
def build_collection(lines):
    try:
        chroma_client.delete_collection("upload_docs")
    except Exception:
        pass
    col = chroma_client.create_collection("upload_docs")
    ids, embeddings, documents = [], [], []
    for i, line in enumerate(lines):
        ids.append(str(i))
        embeddings.append(embed_model.encode(line).tolist())
        documents.append(line)
    col.add(ids=ids, embeddings=embeddings, documents=documents)
    return col

def retrieve(query, collection, n=5):
    vec = embed_model.encode(query).tolist()
    res = collection.query(query_embeddings=[vec], n_results=min(n, len(collection.get()["ids"])))
    return res["documents"][0]

# ── LLM CALLS ────────────────────────────────────────────────
groq_client = None

def get_groq():
    global groq_client
    if groq_client is None:
        groq_client = Groq(api_key=GROQ_API_KEY)
    return groq_client

def llm(system, user, max_tokens=512):
    resp = get_groq().chat.completions.create(
        model=GROQ_MODEL,
        messages=[{"role": "system", "content": system},
                  {"role": "user",   "content": user}],
        max_tokens=max_tokens,
        temperature=0.3,
    )
    return resp.choices[0].message.content.strip()

def detect_subject(text):
    """Ask LLM to identify the subject from the document content."""
    sample = text[:3000]  # first 3000 chars is enough
    return llm(
        "You are an academic subject classifier. Given a document excerpt, "
        "identify the exact subject/course name in 2-5 words. "
        "Reply with ONLY the subject name, nothing else.",
        f"Document excerpt:\n{sample}"
    )

def extract_topics(text, subject):
    """Ask LLM to extract key topics from the document."""
    sample = text[:4000]
    raw = llm(
        f"You are an expert in {subject}. Extract the main topics/chapters from this document. "
        "Return ONLY a numbered list of topic names, one per line, no explanations.",
        f"Document:\n{sample}"
    )
    topics = []
    for line in raw.split("\n"):
        line = line.strip().lstrip("0123456789.-) ").strip()
        if line:
            topics.append(line)
    return topics[:10]

def generate_important_question(topic, collection, subject):
    """RAG: retrieve relevant content, then ask LLM for the most important exam question + answer."""
    docs = retrieve(topic, collection, n=4)
    context = "\n".join(docs)
    return llm(
        f"You are an expert {subject} exam coach. "
        "Based on the provided content, generate the single most important exam question "
        "for this topic, then give a concise model answer. "
        "Format:\nQ: <question>\nA: <answer>",
        f"Topic: {topic}\n\nRelevant content:\n{context}",
        max_tokens=600
    )

def generate_summary(subject, topics):
    topic_list = "\n".join(f"- {t}" for t in topics)
    return llm(
        f"You are a {subject} exam preparation expert. "
        "Given these topics from a question paper, write a short study priority guide "
        "(3-5 sentences) telling students what to focus on most.",
        f"Topics found:\n{topic_list}",
        max_tokens=300
    )

def llm_ask(query, subject):
    return llm(
        f"You are an expert exam preparation assistant{' for ' + subject if subject else ''}. "
        "Give a clear, structured answer with examples. Use bullet points where helpful.",
        query,
        max_tokens=1024
    )

# ── MAIN ─────────────────────────────────────────────────────
if __name__ == "__main__":
    init_db()
    print("\n[SERVER] Running at http://localhost:5000\n")
    app.run(debug=False, port=5000)
